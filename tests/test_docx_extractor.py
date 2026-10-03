"""Automated tests for DOCX resume extraction, error handling, and API integration."""

import io
import zipfile
import pytest
from fastapi.testclient import TestClient

import docx
from app.core import settings
from app.main import app
from app.schemas.resume import ResumeExtractionResponse
from app.services.docx_extractor import (
    DocxEmptyError,
    DocxEncryptedError,
    DocxExtractorService,
    DocxInvalidSignatureError,
    DocxMalformedError,
    DocxPageLimitExceededError,
    DocxUnsupportedLegacyDocError,
    docx_extractor_service,
)
from app.services.contact_extractor import extract_contact_info
from app.services.feedback_analyzer import feedback_analyzer_service
from app.services.job_matcher import job_matcher_service
from app.services.resume_parser import resume_parser_service
from app.services.section_detector import detect_sections
from app.services.skill_extractor import extract_skills

client = TestClient(app)


# ---------------------------------------------------------------------------
# Synthetic DOCX Generation Helpers
# ---------------------------------------------------------------------------

def create_in_memory_docx(
    paragraphs: list[str] = None,
    headings: list[tuple[str, int]] = None,
    bullet_lists: list[str] = None,
    table_data: list[list[str]] = None,
    page_break_count: int = 0,
) -> bytes:
    """Generate in-memory valid Word (.docx) document bytes with specified content."""
    doc = docx.Document()

    if headings:
        for text, level in headings:
            doc.add_heading(text, level=level)

    if paragraphs:
        for p in paragraphs:
            doc.add_paragraph(p)

    if bullet_lists:
        for b in bullet_lists:
            doc.add_paragraph(b, style="List Bullet")

    if table_data:
        num_rows = len(table_data)
        num_cols = max(len(row) for row in table_data) if num_rows > 0 else 0
        if num_rows > 0 and num_cols > 0:
            tbl = doc.add_table(rows=num_rows, cols=num_cols)
            for r_idx, row in enumerate(table_data):
                for c_idx, val in enumerate(row):
                    tbl.cell(r_idx, c_idx).text = val

    for _ in range(page_break_count):
        doc.add_page_break()
        doc.add_paragraph("Additional page content following page break.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SAMPLE_DOCX_RESUME_PARAGRAPHS = [
    "Jane Doe",
    "jane.doe@example.com | +1 (555) 234-5678 | https://github.com/janedoe | https://linkedin.com/in/janedoe",
    "EDUCATION",
    "Bachelor of Science in Computer Science, UC Berkeley (2020 - 2024)",
    "PROFESSIONAL EXPERIENCE",
    "Software Engineer Intern | Tech Innovations Inc. (June 2023 - August 2023)",
    "- Optimized PostgreSQL database indexes, reducing query latency by 45% across 12 high-traffic endpoints.",
    "- Built automated data pipeline with FastAPI and Docker processing over 50,000 requests per day.",
    "- Improved frontend load speed by 2x, resulting in an estimated $5,000 annual cloud infrastructure savings.",
    "PROJECTS",
    "Student Resume Analyzer | Python, FastAPI, React",
    "- Developed an open-source resume review platform serving 1,500 users in university hackathon.",
    "- Achieved 98% test coverage across 65 unit and integration tests using pytest.",
    "TECHNICAL SKILLS",
    "Languages: Python, JavaScript, TypeScript, SQL, C++",
    "Frameworks & Tools: FastAPI, React, Docker, Git, PostgreSQL",
]


# ---------------------------------------------------------------------------
# 1. Unit Tests for DocxExtractorService
# ---------------------------------------------------------------------------

def test_docx_service_normal_extraction():
    """Verify that a standard DOCX document is extracted with complete text and diagnostics."""
    raw_docx = create_in_memory_docx(paragraphs=SAMPLE_DOCX_RESUME_PARAGRAPHS)
    res = docx_extractor_service.extract_text_from_bytes(raw_docx, filename="resume.docx")

    assert isinstance(res, ResumeExtractionResponse)
    assert res.filename == "resume.docx"
    assert res.page_count >= 1
    assert res.character_count > 200
    assert res.word_count > 40
    assert res.parser_status == "success"
    assert res.parser_engine == "python-docx"
    assert res.is_scanned_or_image_based is False
    assert len(res.warnings) == 0
    assert "jane.doe@example.com" in res.extracted_text
    assert "FastAPI" in res.extracted_text


def test_docx_service_multi_page_detection():
    """Verify that explicit page breaks increment page count and populate per-page metadata."""
    raw_docx = create_in_memory_docx(
        paragraphs=["Page 1 content here."],
        page_break_count=2,
    )
    res = docx_extractor_service.extract_text_from_bytes(raw_docx, filename="multipage.docx")

    assert res.page_count == 3
    assert len(res.pages) == 3
    assert res.pages[0].page_number == 1
    assert res.pages[1].page_number == 2
    assert res.pages[2].page_number == 3
    assert all(p.has_text for p in res.pages)


def test_docx_page_count_explicit_breaks_vs_unbroken_paragraphs():
    """Regression test confirming DOCX page counting reflects explicit XML breaks, not print layout.

    Word documents do not encode visual page boundaries unless a rendering engine lays them out.
    A multi-paragraph document with 15 long paragraphs but zero explicit <w:br w:type='page'>
    breaks correctly evaluates to page_count=1, while adding explicit page breaks increments page_count.
    """
    # 1. 15 paragraphs without explicit breaks -> 1 page segment
    paragraphs = [f"Paragraph {i}: Detailed candidate work description and engineering bullets." for i in range(15)]
    doc_without_breaks = create_in_memory_docx(paragraphs=paragraphs, page_break_count=0)
    res_single = docx_extractor_service.extract_text_from_bytes(doc_without_breaks, filename="continuous.docx")
    assert res_single.page_count == 1
    assert len(res_single.pages) == 1

    # 2. Document with 1 explicit break -> 2 page segments
    doc_with_break = create_in_memory_docx(paragraphs=["First section"], page_break_count=1)
    res_multi = docx_extractor_service.extract_text_from_bytes(doc_with_break, filename="broken.docx")
    assert res_multi.page_count == 2
    assert len(res_multi.pages) == 2



def test_docx_service_bullets_and_tables():
    """Verify that bullet lists and table content are properly extracted in document order."""
    raw_docx = create_in_memory_docx(
        headings=[("Skills Summary", 1)],
        paragraphs=["Summary paragraph."],
        bullet_lists=[
            "Architected scalable backend processing 10,000 queries per second.",
            "Decreased memory footprint by 35% through model quantization.",
        ],
        table_data=[
            ["Language", "Python", "JavaScript"],
            ["Database", "PostgreSQL", "MongoDB"],
        ],
    )
    res = docx_extractor_service.extract_text_from_bytes(raw_docx, filename="structured.docx")

    assert "- Architected scalable backend" in res.extracted_text
    assert "- Decreased memory footprint" in res.extracted_text
    assert "Python | JavaScript" in res.extracted_text or "Language | Python" in res.extracted_text


def test_docx_service_empty_bytes_rejected():
    """Verify that 0-byte payload raises DocxEmptyError."""
    with pytest.raises(DocxEmptyError) as exc_info:
        docx_extractor_service.extract_text_from_bytes(b"", filename="empty.docx")
    assert "empty" in str(exc_info.value).lower()


def test_docx_service_empty_text_warning():
    """Verify that a valid DOCX with no readable text triggers the scanned/empty warning."""
    empty_doc = docx.Document()
    buf = io.BytesIO()
    empty_doc.save(buf)

    res = docx_extractor_service.extract_text_from_bytes(buf.getvalue(), filename="blank.docx")
    assert res.character_count == 0
    assert res.is_scanned_or_image_based is True
    assert res.parser_status == "warning_scanned_or_empty"
    assert any("empty or image-based" in w.lower() for w in res.warnings)


def test_docx_service_invalid_signature_rejected():
    """Verify that non-ZIP bytes masquerading as DOCX raise DocxInvalidSignatureError."""
    fake_bytes = b"Just plain text without the PK zip header."
    with pytest.raises(DocxInvalidSignatureError) as exc_info:
        docx_extractor_service.extract_text_from_bytes(fake_bytes, filename="fake.docx")
    assert "signature" in str(exc_info.value).lower()


def test_docx_service_legacy_doc_rejected_with_conversion_message():
    """Verify that legacy binary .doc (OLE2 magic bytes) raises DocxUnsupportedLegacyDocError."""
    # Standard OLE2 compound file magic bytes for legacy Word 97-2003 .doc
    legacy_ole_bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
    with pytest.raises(DocxUnsupportedLegacyDocError) as exc_info:
        docx_extractor_service.extract_text_from_bytes(legacy_ole_bytes, filename="legacy.docx")
    msg = str(exc_info.value).lower()
    assert "legacy binary word" in msg or ".doc" in msg
    assert "convert" in msg or "docx" in msg


def test_docx_service_malformed_zip_rejected():
    """Verify that truncated or damaged ZIP packages raise DocxMalformedError."""
    broken_zip = b"PK\x03\x04\x14\x00\x00\x00\x08\x00truncated_corrupt_data"
    with pytest.raises(DocxMalformedError) as exc_info:
        docx_extractor_service.extract_text_from_bytes(broken_zip, filename="broken.docx")
    assert "malformed" in str(exc_info.value).lower() or "damaged" in str(exc_info.value).lower()


def test_docx_service_encrypted_package_rejected():
    """Verify that encrypted Office documents raise DocxEncryptedError."""
    # OLE2 file with EncryptedPackage stream indicator
    encrypted_ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100 + b"EncryptedPackage" + b"\x00" * 400
    with pytest.raises(DocxEncryptedError) as exc_info:
        docx_extractor_service.extract_text_from_bytes(encrypted_ole, filename="encrypted.docx")
    assert "encrypted" in str(exc_info.value).lower() or "password" in str(exc_info.value).lower()


def test_docx_service_page_limit_enforced():
    """Verify that exceeding MAX_PAGE_COUNT raises DocxPageLimitExceededError."""
    service = DocxExtractorService(max_page_count=3)
    excessive_docx = create_in_memory_docx(
        paragraphs=["Page 1 content"],
        page_break_count=4,  # Results in 5 pages > 3
    )
    with pytest.raises(DocxPageLimitExceededError) as exc_info:
        service.extract_text_from_bytes(excessive_docx, filename="too_long.docx")
    assert "page limit" in str(exc_info.value).lower()


def test_docx_service_low_text_warning():
    """Verify that extracted text under minimum warning threshold produces warning_low_text."""
    service = DocxExtractorService(min_extracted_chars_warning=150, min_page_chars_scanned=5)
    short_docx = create_in_memory_docx(paragraphs=["Short note with 30 characters."])
    res = service.extract_text_from_bytes(short_docx, filename="short.docx")
    assert res.parser_status == "warning_low_text"
    assert any("unusually low" in w.lower() for w in res.warnings)


# ---------------------------------------------------------------------------
# 2. Downstream Pipeline Integration with DOCX Extracted Text
# ---------------------------------------------------------------------------

def test_docx_downstream_parser_and_feedback_pipeline():
    """Verify that text extracted from DOCX flows seamlessly into sections, skills, metrics, and feedback."""
    raw_docx = create_in_memory_docx(paragraphs=SAMPLE_DOCX_RESUME_PARAGRAPHS)
    res = docx_extractor_service.extract_text_from_bytes(raw_docx, filename="jane_doe.docx")
    extracted_text = res.extracted_text

    # 1. Section detection
    sections, _ = detect_sections(extracted_text)
    assert "education" in sections
    assert "experience" in sections
    assert "projects" in sections
    assert "skills" in sections

    # 2. Contact extraction
    contact = extract_contact_info(extracted_text)
    assert contact.email == "jane.doe@example.com"
    assert contact.phone is not None
    assert "github.com/janedoe" in (contact.github_url or "")
    assert "linkedin.com/in/janedoe" in (contact.linkedin_url or "")

    # 3. Skill extraction
    skills, skills_by_cat = extract_skills(extracted_text)
    skill_names = [s.name for s in skills]
    assert "Python" in skill_names
    assert "FastAPI" in skill_names
    assert "Docker" in skill_names
    assert "PostgreSQL" in skill_names

    # 4. Job matching
    jd_text = "Looking for a Python software engineer with experience in FastAPI, Docker, and PostgreSQL."
    match_res = job_matcher_service.match(extracted_text, jd_text, resume_filename="jane_doe.docx")
    assert match_res.text_similarity > 0.10
    assert match_res.skill_overlap_ratio >= 0.75
    assert "Python" in match_res.matched_skills

    # 5. Feedback rubric
    fb_res = feedback_analyzer_service.generate_feedback(extracted_text, resume_filename="jane_doe.docx")
    assert fb_res.total_score >= 80.0
    assert fb_res.index_label in ["Exemplary", "Proficient"]
    assert len(fb_res.quantified_metrics_detected) >= 3


# ---------------------------------------------------------------------------
# 3. HTTP Endpoint Integration Tests (POST /api/v1/resumes/extract)
# ---------------------------------------------------------------------------

def test_api_extract_valid_docx():
    """Verify that uploading a valid DOCX file returns HTTP 200 with schema-conforming response."""
    raw_docx = create_in_memory_docx(paragraphs=SAMPLE_DOCX_RESUME_PARAGRAPHS)
    files = {
        "file": (
            "student_resume.docx",
            raw_docx,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes/extract", files=files)
    assert response.status_code == 200

    data = response.json()
    validated = ResumeExtractionResponse(**data)
    assert validated.filename == "student_resume.docx"
    assert validated.parser_engine == "python-docx"
    assert "jane.doe@example.com" in validated.extracted_text


def test_api_extract_legacy_doc_rejected():
    """Verify that uploading a .doc file returns HTTP 400 with actionable conversion guidance."""
    files = {"file": ("resume.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "application/msword")}
    response = client.post("/api/v1/resumes/extract", files=files)
    assert response.status_code == 400
    detail = response.json()["detail"].lower()
    assert "legacy binary word" in detail or ".doc" in detail
    assert "convert" in detail or ".docx" in detail


def test_api_extract_doc_renamed_to_docx_rejected():
    """Verify that uploading a binary .doc renamed as .docx is trapped by magic byte signature check."""
    legacy_bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100
    files = {
        "file": (
            "fake_legacy.docx",
            legacy_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes/extract", files=files)
    assert response.status_code == 400
    detail = response.json()["detail"].lower()
    assert "legacy" in detail or "convert" in detail


def test_api_extract_corrupted_docx_rejected():
    """Verify that corrupted DOCX upload returns HTTP 422 Unprocessable Content."""
    files = {
        "file": (
            "corrupted.docx",
            b"PK\x03\x04\x00\x00truncated_corrupt_zip_package",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes/extract", files=files)
    assert response.status_code == 422


def test_api_extract_oversized_docx_rejected():
    """Verify that DOCX exceeding MAX_UPLOAD_SIZE_MB (5 MB) returns HTTP 413."""
    oversized_bytes = b"PK\x03\x04" + b"X" * (5 * 1024 * 1024 + 10)
    files = {
        "file": (
            "huge.docx",
            oversized_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes/extract", files=files)
    assert response.status_code == 413
